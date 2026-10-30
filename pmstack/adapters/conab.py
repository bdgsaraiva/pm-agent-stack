"""Conab (Brazil) grain crop surveys, as first published: one row per crop per monthly survey (levantamento).

Conab publishes about 12 surveys per crop year on gov.br/conab, each with a dated page and a "previsão de safra por
produto" workbook. Every survey is kept as its own vintage, like the STEO adapter:

  conab:<crop>:production:<crop year>@<publication date>   Brazil production, 1,000 t, dated at publication

so a reader as of D uses the latest survey published by D, and the revision vs the previous survey is exact.
Crops: soja, milho_total, milho_2a (safrinha), trigo (labelled by harvest year: see wheat_year). Workbooks are .xlsx from about 2021/22; older surveys
published as .xls are skipped and listed in the errors (no xlrd dependency). Brazilian government open data:
cite "Conab, Acompanhamento da Safra Brasileira de Grãos".
"""
from __future__ import annotations

import io
import re
from datetime import date

import pandas as pd

from . import ai_common as ac, base

KIND = "obs"
BASE = "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/boletim-da-safra-de-graos"
_SLUG = re.compile(r"boletim-da-safra-de-graos/(\d{1,2}o-levantamento-safra-20\d\d-(?:20)?\d\d)")
_PUB = re.compile(r"(\d{2})/(\d{2})/(20\d\d) \d{2}h\d{2}")
CROPS = {"SOJA": "soja", "MILHO TOTAL": "milho_total", "MILHO 2": "milho_2a", "TRIGO": "trigo"}


def _norm(s) -> str:
    s = str(s or "").upper().strip()
    return s.replace("ª", "").replace("ª", "").replace("º", "")


def parse_workbook(content: bytes) -> dict[str, tuple[str, float]]:
    """{crop: (crop year label of the current estimate, production in 1,000 t)} from the production sheet."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    name = next((s for s in wb.sheetnames if s.lower().startswith("produ") and "tividade" not in s.lower() and "brasil" in s.lower()), None)
    if name is None:
        return {}
    rows = list(wb[name].iter_rows(values_only=True))
    out = {}
    cur_col, year_lab = None, None
    for i, r in enumerate(rows):
        if r and r[1] == "SAFRAS" and i + 2 < len(rows):          # a table header: years on the next row, (a)/(b)/(c) below
            years, marks = rows[i + 1], rows[i + 2]
            cur_col = next((j for j, m in enumerate(marks) if m and "(c)" in str(m)), None)
            labels = [str(y).strip() for y in years if y and re.match(r"\d", str(y).strip())]
            year_lab = labels[-1] if labels else None
            continue
        if cur_col is None or not r or not r[0]:
            continue
        key = _norm(r[0])
        for prefix, crop in CROPS.items():
            hit = key == prefix if prefix in ("SOJA", "TRIGO", "MILHO TOTAL") else key.startswith(prefix) and "SAFRA" in key
            if hit and crop not in out and isinstance(r[cur_col], (int, float)):
                out[crop] = (year_lab, float(r[cur_col]))
    return out


MONTHS = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}


def pub_date(page_date: date | None, xlsx: str | None) -> date | None:
    """The page's "Publicado em" date, unless it falls outside the survey month named in the file (old pages were
    re-published when gov.br migrated): then the last day of that month, a conservative availability date."""
    m = re.search(r"-(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)-(20\d\d)", xlsx or "")
    if not m:
        return page_date
    y, mo = int(m.group(2)), MONTHS[m.group(1)]
    if page_date and page_date.year == y and page_date.month == mo:
        return page_date
    import calendar
    return date(y, mo, calendar.monthrange(y, mo)[1])


_LINK = re.compile(r'href="([^"]*(?:previs[^"]*por[_-]produto|tabela-de-dados-producao)[^"]*?)(?:/view)?"')


def wheat_year(pub: date) -> int:
    """Brazilian wheat is harvested September-December; Conab's surveys switch to the next crop in February. So a
    survey from February to December estimates that year's harvest, a January survey the previous year's."""
    return pub.year if pub.month >= 2 else pub.year - 1


def survey(slug: str, cfg: dict) -> tuple[date | None, str | None]:
    page = ac.get(f"{BASE}/{slug}", cfg, headers={"User-Agent": ac.BROWSER_UA}).text
    if not _LINK.search(page) and f"{slug}/{slug}" in page:     # some surveys are a folder holding a same-named page
        page = ac.get(f"{BASE}/{slug}/{slug}", cfg, headers={"User-Agent": ac.BROWSER_UA}).text + page
    m = _PUB.search(page)
    pd_ = date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None
    link = _LINK.search(page)
    xlsx = link.group(1) if link else None
    return pub_date(pd_, xlsx), xlsx


def fetch(cfg: dict, since=None) -> pd.DataFrame:
    fetched = base.now_utc()
    budget = ac.Budget(cfg, 240)
    pages = [0] if since is not None else list(range(0, 30 * int(cfg.get("listing_pages", 4)), 30))
    slugs: list[str] = []
    for b in pages:
        html = ac.get(f"{BASE}?b_start:int={b}", cfg, headers={"User-Agent": ac.BROWSER_UA}).text
        for s in _SLUG.findall(html):
            if s not in slugs:
                slugs.append(s)
    rows, errors = [], []
    for slug in slugs:
        if budget.out():
            errors.append(f"time budget: {len(slugs) - slugs.index(slug)} surveys left for next run")
            break
        try:
            pub, xlsx = survey(slug, cfg)
            if not xlsx or not pub:
                errors.append(f"{slug}: no workbook link on the survey page")
                continue
            content = ac.get(xlsx, cfg, headers={"User-Agent": ac.BROWSER_UA}).content
            if not content.startswith(b"PK"):                  # .xls (OLE2) or an HTML page: not read
                errors.append(f"{slug}: workbook is not .xlsx (older .xls not read)")
                continue
            vals = parse_workbook(content)
        except Exception as e:  # noqa: BLE001
            errors.append(f"{slug}: {type(e).__name__}: {e}"[:200])
            continue
        if not vals:
            errors.append(f"{slug}: production sheet not found")
        for crop, (year, v) in vals.items():
            if crop == "trigo":                                 # review S4: the workbook's winter-crop year label moves to
                year = str(wheat_year(pub))                     # the next year in October while the value is still this
            yr = (year or "").replace("/", "-")                 # year's harvest; label by harvest year instead
            rows.append(ac.row("conab", f"conab:{crop}:production:{yr}@{pub.isoformat()}",
                               f"Conab {crop} production {year}, survey published {pub} ({slug})", pub, v, "kt", "survey",
                               f"{BASE}/{slug}", fetched))
    if not rows:
        raise RuntimeError("; ".join(errors) or "no Conab survey parsed")
    return ac.frame(rows, errors)
