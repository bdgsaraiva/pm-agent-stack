"""EIA Short-Term Energy Outlook as first published: one workbook per monthly issue (free, no key; public domain).

Each issue is stored under its own series ids, `steo:<SERIES>@<YYYY-MM>` (issue month), with one row per target
month, so EIA's view at any past date and its revisions are recoverable (vintages back to the archive's first xlsx).
STEO is "EIA's view", never the market's expectation: EIA's stock paths barely beat a seasonal-naive path
(agent 9 v2 research, notes-A14), so no surprise is ever measured against it.

Archive: https://www.eia.gov/outlooks/steo/archives/{mon}{yy}_base.xlsx (older issues are .xls, not read here);
current issue: https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx. Released monthly around noon ET (release schedule).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone

import pandas as pd

from . import base

KIND = "obs"
ARCHIVE = "https://www.eia.gov/outlooks/steo/archives/{m}{y}_base.xlsx"
CURRENT = "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx"
MON = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
MONTH_ABBR = {m.capitalize(): i + 1 for i, m in enumerate(MON)}


def parse(content: bytes, sheets: list[str], series: set[str]) -> tuple[date | None, list[tuple[str, date, float]]]:
    """(forecast date, [(series, target month, value)]) from one STEO workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    fdate, rows, seen = None, [], set()
    series = {x.upper() for x in series}
    for name in sheets:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        years, months, cols = None, None, []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 2:
                years = list(row)
            elif i == 3:
                months = list(row)
                if fdate is None and row[0]:
                    try:
                        fdate = datetime.strptime(re.sub(r"\s+", " ", str(row[0]).split(",", 1)[1].strip()), "%B %d, %Y").date()
                    except (ValueError, IndexError):
                        pass
                y = None
                for j in range(2, len(months)):
                    if years[j] not in (None, ""):
                        try:
                            y = int(float(years[j]))
                        except (TypeError, ValueError):
                            pass
                    m = MONTH_ABBR.get(str(months[j]).strip()[:3]) if months[j] else None
                    cols.append((j, date(y, m, 1)) if y and m else (j, None))
            elif i > 3 and row and row[0] and str(row[0]).strip().upper() in series:
                sid = str(row[0]).strip().upper()               # world table ids are lower case (papr_world)
                if (name, sid) in seen:                          # papr_world appears twice in 3atab: keep the first
                    continue
                seen.add((name, sid))
                for j, d in cols:
                    if d and j < len(row) and isinstance(row[j], (int, float)):
                        rows.append((sid, d, float(row[j])))
    return fdate, rows


def issues(since: date, today: date) -> list[tuple[str, str]]:
    out, y, m = [], since.year, since.month
    while (y, m) <= (today.year, today.month):
        out.append((f"{y:04d}-{m:02d}", ARCHIVE.format(m=MON[m - 1], y=f"{y % 100:02d}")))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def fetch(cfg: dict, since=None) -> pd.DataFrame:
    import requests
    today = datetime.now(timezone.utc).date()
    series = set(cfg["series"])
    known = cfg.get("_known") or set()
    have = {k.split("@", 1)[1] for k in known if "@" in k}
    import json as _json
    from pathlib import Path as _P
    skip_f = _P(__file__).resolve().parents[2] / "data" / "obs" / "steo" / "_skip.json"
    skip = set(_json.loads(skip_f.read_text(encoding="utf-8"))) if skip_f.exists() else set()
    cur_iss = f"{today.year:04d}-{today.month:02d}"
    have |= {x for x in skip if x != cur_iss}       # archive months with no xlsx are not re-downloaded every run
    todo = [(iss, url) for iss, url in issues(date.fromisoformat(cfg.get("archive_from", "2014-01-01")), today) if iss not in have]
    fetched = base.now_utc()
    frames, errors = [], []
    t0 = datetime.now(timezone.utc)
    for iss, url in sorted(todo, reverse=True):              # newest first
        if (datetime.now(timezone.utc) - t0).total_seconds() > float(cfg.get("budget_s", 150)):
            errors.append(f"budget reached; {iss} and older wait for the next run")
            break
        try:
            r = requests.get(url, timeout=int(cfg.get("timeout_s", 60)), headers={"User-Agent": "pm-agent-stack (personal research)"})
            if r.status_code == 404 and iss == f"{today.year:04d}-{today.month:02d}":
                url = CURRENT                                    # this month's issue is not archived yet
                r = requests.get(url, timeout=int(cfg.get("timeout_s", 60)), headers={"User-Agent": "pm-agent-stack (personal research)"})
            if r.status_code == 404:
                if iss != cur_iss:
                    skip.add(iss)
                continue                                         # no xlsx for that month (older .xls, or not issued yet)
            r.raise_for_status()
            fdate, rows = parse(r.content, cfg["sheets"], series)
        except Exception as e:  # noqa: BLE001
            errors.append(f"{iss}: {e}")
            continue
        if url == CURRENT and fdate and f"{fdate.year:04d}-{fdate.month:02d}" != iss:
            continue                                             # the current file is still last month's issue
        if not rows:
            continue
        df = pd.DataFrame(rows, columns=["sid", "date", "value"])
        frames.append(pd.DataFrame({"source": "steo", "series_id": "steo:" + df["sid"] + "@" + iss,
                                    "name": "STEO " + df["sid"] + f" (issue {iss}, forecast date {fdate})",
                                    "date": df["date"], "value": df["value"], "unit": "steo", "freq": "monthly",
                                    "source_url": url, "fetched_at": fetched}))
    if skip:
        skip_f.parent.mkdir(parents=True, exist_ok=True)
        skip_f.write_text(_json.dumps(sorted(skip)), encoding="utf-8")
    out = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    out.attrs["errors"] = errors
    return out
